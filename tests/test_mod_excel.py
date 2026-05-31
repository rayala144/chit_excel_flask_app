import pytest
from io import BytesIO
from openpyxl import Workbook, load_workbook

from mod_excel import PDF_SHEET_NAMES, prepare_workbook_for_pdf, update_excel


def test_update_excel_writes_grand_total(sample_workbook):
    updated = update_excel(sample_workbook)

    assert updated["Sheet3"]["E43"].value == "GRAND TOTAL"
    assert updated["Sheet3"]["F43"].value is not None
    assert updated["Sheet3"]["F43"].value > 0


def test_update_excel_fills_sum_columns(sample_workbook):
    updated = update_excel(sample_workbook)

    assert updated["Sheet2"]["C3"].value == 10
    assert updated["Sheet2"]["F3"].value == 20


def test_prepare_workbook_for_pdf_keeps_only_pdf_sheets(sample_workbook):
    pdf_bytes = prepare_workbook_for_pdf(sample_workbook)
    workbook = load_workbook(BytesIO(pdf_bytes))

    assert workbook.sheetnames == list(PDF_SHEET_NAMES)


def test_prepare_workbook_for_pdf_missing_sheet_raises():
    workbook = Workbook()
    workbook.active.title = "Sheet1"

    with pytest.raises(ValueError, match="missing sheets required for PDF"):
        prepare_workbook_for_pdf(workbook)
