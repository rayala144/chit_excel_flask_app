import importlib
from io import BytesIO

import pytest
from openpyxl import Workbook

import app as app_module


def reset_supabase_state() -> None:
    app_module._supabase_client = None
    app_module._supabase_init_attempted = False


@pytest.fixture(autouse=True)
def _isolate_supabase(monkeypatch):
    reset_supabase_state()
    monkeypatch.delenv("SUPABASE_URL", raising=False)
    monkeypatch.delenv("SUPABASE_KEY", raising=False)
    monkeypatch.delenv("SUPABASE_SERVICE_KEY", raising=False)
    monkeypatch.delenv("SUPABASE_ANON_KEY", raising=False)
    yield
    reset_supabase_state()


@pytest.fixture
def flask_app():
    app_module.app.config["TESTING"] = True
    return app_module.app


@pytest.fixture
def client(flask_app):
    return flask_app.test_client()


def build_sample_workbook() -> Workbook:
    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = "Sheet1"
    workbook.create_sheet("Sheet2")
    workbook.create_sheet("Sheet3")

    sheet1["B3"] = 10
    sheet1["B4"] = 20
    sheet1["B5"] = 30

    for sheet_name in ("Sheet2", "Sheet3"):
        sheet = workbook[sheet_name]
        for row in range(3, 33 if sheet_name == "Sheet2" else 41):
            sheet[f"B{row}"] = "1"
            sheet[f"E{row}"] = "2"

    return workbook


@pytest.fixture
def sample_workbook():
    return build_sample_workbook()


@pytest.fixture
def sample_xlsx_bytes(sample_workbook):
    buffer = BytesIO()
    sample_workbook.save(buffer)
    return buffer.getvalue()
